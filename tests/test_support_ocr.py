"""Document image extraction checks for the support knowledge base."""

from io import BytesIO
from types import SimpleNamespace

import pytest
from docx import Document as WordDocument
from PIL import Image
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject, NumberObject

from app.services import support_documents


def sample_image(format_name: str = "PNG") -> bytes:
    """Build a small raster image for document fixtures."""
    output = BytesIO()
    Image.new("RGB", (600, 160), "white").save(output, format=format_name)
    return output.getvalue()


def test_image_upload_runs_local_ocr(monkeypatch: pytest.MonkeyPatch) -> None:
    """Image uploads flow through Tesseract and become reviewable source text."""
    calls: list[list[str]] = []

    def fake_run(command: list[str], **kwargs: object) -> SimpleNamespace:
        calls.append(command)
        assert kwargs["input"]
        return SimpleNamespace(returncode=0, stdout="云存档说明".encode())

    monkeypatch.setattr(support_documents.subprocess, "run", fake_run)
    assert support_documents.parse_document_pages("guide.png", sample_image()) == [(None, "云存档说明")]
    assert calls[0][1:3] == ["stdin", "stdout"]


def test_docx_embedded_image_text_is_included(monkeypatch: pytest.MonkeyPatch) -> None:
    """A DOCX screenshot contributes OCR text beside its selectable text."""
    document = WordDocument()
    document.add_paragraph("存档帮助")
    document.add_picture(BytesIO(sample_image()))
    output = BytesIO()
    document.save(output)
    monkeypatch.setattr(support_documents, "_ocr_image", lambda raw: "设置界面中的云存档开关")

    pages = support_documents.parse_document_pages("guide.docx", output.getvalue())

    assert "存档帮助" in pages[0][1]
    assert "[Image text]" in pages[0][1]
    assert "设置界面中的云存档开关" in pages[0][1]


def test_scanned_pdf_keeps_page_number(monkeypatch: pytest.MonkeyPatch) -> None:
    """OCR text from a PDF page keeps its page citation."""
    writer = PdfWriter()
    page = writer.add_blank_page(width=600, height=160)
    image = DecodedStreamObject()
    image.set_data(sample_image("JPEG"))
    image.update(
        {
            NameObject("/Type"): NameObject("/XObject"),
            NameObject("/Subtype"): NameObject("/Image"),
            NameObject("/Width"): NumberObject(600),
            NameObject("/Height"): NumberObject(160),
            NameObject("/ColorSpace"): NameObject("/DeviceRGB"),
            NameObject("/BitsPerComponent"): NumberObject(8),
            NameObject("/Filter"): NameObject("/DCTDecode"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/XObject"): DictionaryObject({NameObject("/Im1"): writer._add_object(image)})}
    )
    content = DecodedStreamObject()
    content.set_data(b"q 600 0 0 160 0 0 cm /Im1 Do Q")
    page[NameObject("/Contents")] = writer._add_object(content)
    output = BytesIO()
    writer.write(output)
    monkeypatch.setattr(support_documents, "_ocr_image", lambda raw: "玩家设置截图")

    assert support_documents.parse_document_pages("scan.pdf", output.getvalue()) == [(1, "玩家设置截图")]
