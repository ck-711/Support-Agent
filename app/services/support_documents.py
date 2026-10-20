"""Parse, embed, index, and retrieve approved support documents."""

import asyncio
import json
import re
import subprocess
import time
from collections import Counter
from dataclasses import dataclass
from datetime import timedelta
from io import BytesIO
from pathlib import Path
from typing import Awaitable, Callable, Sequence

from docx import Document as WordDocument
from langchain_text_splitters import RecursiveCharacterTextSplitter
from openai import AsyncOpenAI
from PIL import Image, UnidentifiedImageError
from pypdf import PdfReader
from sqlalchemy import case, delete, func, literal_column, or_
from sqlmodel import col, select
from tenacity import retry, stop_after_attempt, wait_exponential

from app.core.config import settings
from app.core.logging import logger
from app.models.support import SupportChunk, SupportDocument, SupportDocumentImage, SupportKnowledge, now_utc
from app.services.support_milvus import SupportVector, replace_source, search as search_milvus
from app.services.support_store import session_factory

MAX_UPLOAD_BYTES = 50 * 1024 * 1024
MAX_TEXT_LENGTH = 600_000
MAX_CHUNKS = 800
EMBEDDING_BATCH_SIZE = 10
MAX_PDF_PAGES = 200
SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx", ".png", ".jpg", ".jpeg", ".webp"}
MAX_OCR_IMAGES = 400
MAX_OCR_PIXELS = 20_000_000
MAX_OCR_DOCUMENT_SECONDS = 300
CHINESE_STOPWORDS = {
    "游戏",
    "玩家",
    "可以",
    "如果",
    "然后",
    "这个",
    "那个",
    "进行",
    "通过",
    "使用",
    "需要",
    "可能",
    "如何",
    "什么",
    "一个",
    "以及",
    "但是",
    "因为",
    "所以",
    "之后",
    "之前",
    "现在",
    "问题",
    "内容",
}
LATIN_STOPWORDS = {
    "the",
    "and",
    "for",
    "that",
    "with",
    "this",
    "from",
    "have",
    "your",
    "will",
    "can",
    "are",
    "was",
    "were",
    "when",
    "then",
    "than",
    "not",
    "you",
    "all",
    "any",
    "how",
    "what",
    "which",
    "into",
    "their",
    "game",
    "games",
    "player",
    "players",
    "please",
    "using",
    "used",
    "also",
    "more",
    "most",
    "only",
    "some",
    "los",
    "las",
    "una",
    "uno",
    "unos",
    "unas",
    "para",
    "por",
    "con",
    "como",
    "pero",
    "sobre",
    "este",
    "esta",
    "esto",
    "cuando",
    "desde",
    "puede",
    "pueden",
    "juego",
    "jugador",
    "jugadores",
    "tiene",
    "hacer",
}


@dataclass(frozen=True)
class RetrievedPassage:
    """A published source passage with stable citation URL."""

    id: str
    title: str
    content: str
    url: str
    distance: float
    page_number: int | None = None
    image_id: str | None = None


async def _image_embeddings(document_id: str | None) -> list[tuple[str, list[float]]]:
    """Rebuild reviewed image vectors when a document changes embedding model."""
    if document_id is None:
        return []
    async with session_factory() as db:
        result = await db.exec(
            select(SupportDocumentImage)
            .where(col(SupportDocumentImage.document_id) == document_id, col(SupportDocumentImage.description) != "")
        )
        images = list(result.all())
    if not images:
        return []
    vectors = await embed_texts([image.description for image in images])
    return [(image.id, vector) for image, vector in zip(images, vectors, strict=True)]


def _ocr_image(raw: bytes) -> str:
    """Recognize one image with a bounded local Tesseract process."""
    try:
        with Image.open(BytesIO(raw)) as image:
            if image.width * image.height > MAX_OCR_PIXELS:
                raise ValueError("Image exceeds the 20-megapixel OCR limit")
            if image.width < 80 or image.height < 40:
                return ""
            output = BytesIO()
            image.convert("RGB").save(output, format="PNG")
    except (UnidentifiedImageError, OSError) as exc:
        raise ValueError("An embedded image could not be decoded") from exc
    try:
        result = subprocess.run(
            [settings.SUPPORT_OCR_COMMAND, "stdin", "stdout", "-l", settings.SUPPORT_OCR_LANGUAGES],
            input=output.getvalue(),
            capture_output=True,
            check=False,
            timeout=20,
        )
    except FileNotFoundError as exc:
        raise ValueError("OCR unavailable: install Tesseract or configure SUPPORT_OCR_COMMAND") from exc
    except subprocess.TimeoutExpired as exc:
        raise ValueError("OCR timed out while reading an image") from exc
    if result.returncode:
        raise ValueError("OCR failed: check Tesseract language packs and SUPPORT_OCR_LANGUAGES")
    return result.stdout.decode("utf-8", errors="replace").strip()


def _append_image_text(text: str, image_text: str) -> str:
    """Avoid adding an OCR result already present in the selectable text layer."""
    if not image_text:
        return text
    normalized_source = re.sub(r"\s+", "", text).casefold()
    normalized_ocr = re.sub(r"\s+", "", image_text).casefold()
    if normalized_ocr in normalized_source:
        return text
    return f"{text}\n\n[Image text]\n{image_text}" if text else image_text


def parse_document_pages(filename: str, raw: bytes) -> list[tuple[int | None, str]]:
    """Extract text while preserving PDF page provenance."""
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError("Supported formats: TXT, Markdown, PDF, DOCX, PNG, JPG, WEBP")
    if not raw or len(raw) > MAX_UPLOAD_BYTES:
        raise ValueError("Document must be between 1 byte and 50 MiB")
    pages: list[tuple[int | None, str]]
    try:
        if suffix in {".txt", ".md"}:
            pages = [(None, raw.decode("utf-8-sig"))]
        elif suffix == ".pdf":
            reader = PdfReader(BytesIO(raw))
            if len(reader.pages) > MAX_PDF_PAGES:
                raise ValueError("PDF exceeds the 200-page limit")
            pages = []
            image_count = 0
            deadline = time.monotonic() + MAX_OCR_DOCUMENT_SECONDS
            for number, page in enumerate(reader.pages, start=1):
                text = page.extract_text() or ""
                images = page.images
                image_count += len(images)
                if image_count > MAX_OCR_IMAGES:
                    raise ValueError("Document exceeds the 400-image OCR limit")
                for image in images:
                    if time.monotonic() >= deadline:
                        raise ValueError("Document OCR exceeds the five-minute limit")
                    text = _append_image_text(text, _ocr_image(image.data))
                pages.append((number, text))
        elif suffix in {".png", ".jpg", ".jpeg", ".webp"}:
            pages = [(None, _ocr_image(raw))]
        else:
            document = WordDocument(BytesIO(raw))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            for table in document.tables:
                text += "\n" + "\n".join(" | ".join(cell.text for cell in row.cells) for row in table.rows)
            images = [part for part in document.part.related_parts.values() if part.content_type.startswith("image/")]
            if len(images) > MAX_OCR_IMAGES:
                raise ValueError("Document exceeds the 400-image OCR limit")
            deadline = time.monotonic() + MAX_OCR_DOCUMENT_SECONDS
            for image in images:
                if time.monotonic() >= deadline:
                    raise ValueError("Document OCR exceeds the five-minute limit")
                text = _append_image_text(text, _ocr_image(image.blob))
            pages = [(None, text)]
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Document could not be parsed") from exc
    # PostgreSQL text columns reject NUL bytes that can appear in PDF text layers.
    pages = [
        (number, re.sub(r" *\x00+ *", " ", text).strip())
        for number, text in pages
        if text.strip()
    ]
    text = "\n\n".join(content for _, content in pages)
    if not text or len(text) > MAX_TEXT_LENGTH:
        raise ValueError("Document must contain extractable text under 600,000 characters")
    return pages


def parse_document(filename: str, raw: bytes) -> str:
    """Extract plain text for preview and existing callers."""
    return "\n\n".join(content for _, content in parse_document_pages(filename, raw))


def split_document(content: str) -> list[str]:
    """Split source text into overlapping retrieval passages."""
    splitter = RecursiveCharacterTextSplitter(chunk_size=1200, chunk_overlap=180)
    chunks = [chunk.strip() for chunk in splitter.split_text(content) if chunk.strip()]
    if not chunks or len(chunks) > MAX_CHUNKS:
        raise ValueError("Document exceeds the 800-chunk index limit")
    return chunks


def split_document_pages(pages: Sequence[tuple[int | None, str]]) -> list[tuple[int | None, str]]:
    """Split within pages so citations never cross a PDF page boundary."""
    chunks = [(page, chunk) for page, content in pages for chunk in split_document(content)]
    if len(chunks) > MAX_CHUNKS:
        raise ValueError("Document exceeds the 800-chunk index limit")
    return chunks


def extract_keyword_candidates(content: str, limit: int = 10) -> list[str]:
    """Suggest useful English terms and Chinese character n-grams for review."""
    latin_terms = (
        term
        for term in re.findall(r"[a-záéíóúüñ][a-záéíóúüñ0-9_-]{2,}", content.casefold())
        if term not in LATIN_STOPWORDS
    )
    terms = Counter(latin_terms)
    for phrase in re.findall(r"[\u3400-\u9fff]{2,}", content):
        for size in (2, 3):
            for start in range(max(0, len(phrase) - size + 1)):
                candidate = phrase[start : start + size]
                if candidate not in CHINESE_STOPWORDS:
                    terms[candidate] += 1
    return [term for term, _ in sorted(terms.items(), key=lambda item: (-item[1], -len(item[0]), item[0]))[:limit]]


def normalize_keywords(keywords: Sequence[str]) -> list[str]:
    """Normalize staff-entered aliases while preserving their language."""
    return list(dict.fromkeys(value.strip()[:80] for value in keywords if value.strip()))[:30]


def stored_keywords(value: str) -> list[str]:
    """Decode keywords from older or current database rows safely."""
    try:
        parsed = json.loads(value)
    except (TypeError, json.JSONDecodeError):
        return []
    return (
        normalize_keywords(parsed)
        if isinstance(parsed, list) and all(isinstance(item, str) for item in parsed)
        else []
    )


@retry(stop=stop_after_attempt(3), wait=wait_exponential(multiplier=1, min=1, max=8), reraise=True)
async def _embed_batch(client: AsyncOpenAI, texts: list[str]) -> list[list[float]]:
    """Retry transient embedding provider failures with exponential backoff."""
    response = await client.embeddings.create(
        model=settings.SUPPORT_EMBEDDING_MODEL,
        input=texts,
        dimensions=settings.SUPPORT_EMBEDDING_DIMENSIONS,
    )
    vectors = [item.embedding for item in sorted(response.data, key=lambda item: item.index)]
    if len(vectors) != len(texts) or any(len(vector) != 1024 for vector in vectors):
        raise ValueError("Embedding endpoint must return one 1024-dimensional vector per passage")
    return vectors


async def embed_texts(
    texts: list[str], progress: Callable[[int, int], Awaitable[None]] | None = None
) -> list[list[float]]:
    """Use the configured OpenAI-compatible embedding API."""
    if not settings.SUPPORT_EMBEDDING_API_KEY:
        raise ValueError("SUPPORT_EMBEDDING_API_KEY is not configured")
    options = {"api_key": settings.SUPPORT_EMBEDDING_API_KEY, "timeout": 30.0}
    if settings.SUPPORT_EMBEDDING_BASE_URL:
        options["base_url"] = settings.SUPPORT_EMBEDDING_BASE_URL
    async with AsyncOpenAI(**options) as client:  # pyright: ignore[reportArgumentType]
        vectors: list[list[float]] = []
        for start in range(0, len(texts), EMBEDDING_BATCH_SIZE):
            vectors.extend(await _embed_batch(client, texts[start : start + EMBEDDING_BATCH_SIZE]))
            if progress is not None:
                await progress(min(start + EMBEDDING_BATCH_SIZE, len(texts)), len(texts))
        return vectors


async def index_document(
    title: str,
    filename: str,
    mime_type: str,
    raw: bytes,
    source_url: str | None = None,
    document_id: str | None = None,
    *,
    source_type: str = "unknown",
    game_version: str | None = None,
    language: str = "en",
    keywords: Sequence[str] = (),
    progress: Callable[[str, int], Awaitable[None]] | None = None,
) -> SupportDocument:
    """Build an index before atomically replacing a document's passages."""
    if progress is not None:
        await progress("parsing", 35)
    pages = await asyncio.to_thread(parse_document_pages, filename, raw)
    if progress is not None:
        await progress("splitting", 45)
    content = "\n\n".join(text for _, text in pages)
    chunks = split_document_pages(pages)
    document_keywords = normalize_keywords(keywords)
    chunk_keywords = [
        normalize_keywords([*document_keywords, *extract_keyword_candidates(chunk)]) for _, chunk in chunks
    ]
    if progress is not None:
        await progress("embedding", 50)

        async def embedding_progress(done: int, total: int) -> None:
            await progress("embedding", 50 + 40 * done // total)

        vectors = await embed_texts([chunk for _, chunk in chunks], progress=embedding_progress)
    else:
        vectors = await embed_texts([chunk for _, chunk in chunks])
    image_vectors = await _image_embeddings(document_id)
    if progress is not None:
        await progress("saving", 95)
    async with session_factory() as db:
        document = (
            await db.get(SupportDocument, document_id)
            if document_id
            else SupportDocument(
                title=title,
                filename=filename,
                mime_type=mime_type,
                content=content,
                embedding_model=settings.SUPPORT_EMBEDDING_MODEL,
            )
        )
        if document is None:
            raise LookupError("Document not found")
        if document_id:
            await db.exec(delete(SupportChunk).where(col(SupportChunk.document_id) == document_id))
        document.title = title
        document.filename = filename
        document.mime_type = mime_type
        document.content = content
        document.source_type = source_type
        document.game_version = game_version
        document.language = language
        document.keywords_json = json.dumps(document_keywords, ensure_ascii=False)
        document.pages_json = json.dumps([{"page": page, "content": text} for page, text in pages], ensure_ascii=False)
        document.source_url = source_url
        document.status = "draft"
        document.reviewed_by = None
        document.review_due_at = None
        document.chunk_count = len(chunks)
        document.embedding_model = settings.SUPPORT_EMBEDDING_MODEL
        document.updated_at = now_utc()
        db.add(document)
        await db.flush()
        for position, ((page, passage), vector, passage_keywords) in enumerate(
            zip(chunks, vectors, chunk_keywords, strict=True)
        ):
            db.add(
                SupportChunk(
                    document_id=document.id,
                    position=position,
                    page_number=page,
                    content=passage,
                    keywords_json=json.dumps(passage_keywords, ensure_ascii=False),
                    embedding=vector,
                )
            )
        for image_id, vector in image_vectors:
            image = await db.get(SupportDocumentImage, image_id)
            if image is not None:
                image.embedding = vector
                db.add(image)
        await db.commit()
        logger.info("support_document_indexed", document_id=document.id, chunk_count=len(chunks))
        return document


async def reindex_document(document_id: str) -> SupportDocument | None:
    """Rebuild vectors for stored text after changing embedding settings."""
    async with session_factory() as db:
        document = await db.get(SupportDocument, document_id)
        if document is None:
            return None
        existing_chunks = await db.exec(select(SupportChunk).where(col(SupportChunk.document_id) == document_id))
        reviewed_keywords = {
            (chunk.page_number, chunk.content): stored_keywords(chunk.keywords_json) for chunk in existing_chunks.all()
        }
        title, filename, mime_type, content, source_url, pages_json, source_type, game_version, language, keywords = (
            document.title,
            document.filename,
            document.mime_type,
            document.content,
            document.source_url,
            document.pages_json,
            document.source_type,
            document.game_version,
            document.language,
            stored_keywords(document.keywords_json),
        )
    pages = [(item["page"], item["content"]) for item in json.loads(pages_json)] if pages_json else [(None, content)]
    chunks = split_document_pages(pages)
    vectors = await embed_texts([chunk for _, chunk in chunks])
    image_vectors = await _image_embeddings(document_id)
    async with session_factory() as db:
        document = await db.get(SupportDocument, document_id)
        if document is None:
            return None
        await db.exec(delete(SupportChunk).where(col(SupportChunk.document_id) == document_id))
        document.title, document.filename, document.mime_type, document.source_url = (
            title,
            filename,
            mime_type,
            source_url,
        )
        document.source_type, document.game_version, document.language = source_type, game_version, language
        document.keywords_json = json.dumps(keywords, ensure_ascii=False)
        document.status = "draft"
        document.reviewed_by = None
        document.review_due_at = None
        document.chunk_count = len(chunks)
        document.embedding_model = settings.SUPPORT_EMBEDDING_MODEL
        document.updated_at = now_utc()
        db.add(document)
        chunk_keywords = [
            normalize_keywords(
                [*keywords, *(reviewed_keywords.get((page, chunk)) or extract_keyword_candidates(chunk))]
            )
            for page, chunk in chunks
        ]
        for position, ((page, passage), vector, passage_keywords) in enumerate(
            zip(chunks, vectors, chunk_keywords, strict=True)
        ):
            db.add(
                SupportChunk(
                    document_id=document_id,
                    position=position,
                    page_number=page,
                    content=passage,
                    keywords_json=json.dumps(passage_keywords, ensure_ascii=False),
                    embedding=vector,
                )
            )
        for image_id, vector in image_vectors:
            image = await db.get(SupportDocumentImage, image_id)
            if image is not None:
                image.embedding = vector
                db.add(image)
        await db.commit()
        return document


async def list_documents() -> list[SupportDocument]:
    """List indexed documents for staff."""
    async with session_factory() as db:
        result = await db.exec(select(SupportDocument).order_by(col(SupportDocument.updated_at).desc()).limit(200))
        return list(result.all())


async def get_document(document_id: str, published_only: bool = False) -> SupportDocument | None:
    """Read a document, optionally enforcing publication."""
    async with session_factory() as db:
        document = await db.get(SupportDocument, document_id)
        return document if document and (not published_only or document.status == "published") else None


async def list_document_chunks(document_id: str) -> list[SupportChunk]:
    """Return a document's passages in index order for staff review."""
    async with session_factory() as db:
        result = await db.exec(
            select(SupportChunk)
            .where(col(SupportChunk.document_id) == document_id)
            .order_by(col(SupportChunk.position))
        )
        return list(result.all())


async def save_chunk_keywords(document_id: str, chunk_id: str, keywords: Sequence[str]) -> SupportChunk | None:
    """Persist staff-reviewed aliases for one indexed passage."""
    async with session_factory() as db:
        chunk = await db.get(SupportChunk, chunk_id)
        if chunk is None or chunk.document_id != document_id:
            return None
        chunk.keywords_json = json.dumps(normalize_keywords(keywords), ensure_ascii=False)
        db.add(chunk)
        await db.commit()
        return chunk


def document_page_text(document: SupportDocument, page_number: int | None) -> str | None:
    """Return the extracted source page referenced by a citation."""
    if page_number is None:
        return document.content
    if not document.pages_json:
        return None
    for item in json.loads(document.pages_json):
        if item["page"] == page_number:
            return str(item["content"])
    return None


async def publish_document(document_id: str, staff_id: int) -> SupportDocument | None:
    """Make indexed passages available to the support agent."""
    async with session_factory() as db:
        document = await db.get(SupportDocument, document_id)
        if document is None or document.chunk_count == 0:
            return None
        images = await db.exec(select(SupportDocumentImage).where(col(SupportDocumentImage.document_id) == document_id))
        image_rows = list(images.all())
        if any(not image.description.strip() for image in image_rows):
            raise ValueError("Every image needs a reviewed description before publication")
        chunks = await db.exec(select(SupportChunk).where(col(SupportChunk.document_id) == document_id))
        published_at = now_utc()
        records = await document_vectors(document, list(chunks.all()), image_rows, published_at.isoformat())
        await replace_source(document.id, records)
        document.status = "published"
        document.reviewed_by = staff_id
        document.updated_at = published_at
        document.review_due_at = document.updated_at + timedelta(days=90)
        db.add(document)
        await db.commit()
        return document


async def delete_document(document_id: str) -> bool:
    """Remove a document and its vectors."""
    async with session_factory() as db:
        document = await db.get(SupportDocument, document_id)
        if document is None:
            return False
        await db.exec(delete(SupportChunk).where(col(SupportChunk.document_id) == document_id))
        await db.delete(document)
        await db.commit()
        try:
            await replace_source(document_id, [])
        except Exception:
            logger.exception("support_milvus_source_delete_failed", document_id=document_id)
        return True


async def article_vectors(article: SupportKnowledge) -> list[SupportVector]:
    """Embed a reviewed article as passages in the shared Milvus collection."""
    passages = split_document_pages([(None, article.content)])
    texts = [f"{article.title}\n{passage}" for _, passage in passages]
    vectors = await embed_texts(texts)
    return [
        SupportVector(
            source_id=article.id, item_id=f"{article.id}:{position}", source_type="article",
            revision=str(article.version), title=article.title, content=passage,
            url=article.source_url, embedding=vector,
        )
        for position, ((_, passage), vector) in enumerate(zip(passages, vectors, strict=True))
    ]


async def document_vectors(
    document: SupportDocument, chunks: list[SupportChunk], images: list[SupportDocumentImage], revision: str
) -> list[SupportVector]:
    """Use stored vectors when their model matches, otherwise embed the existing source text."""
    current_model = document.embedding_model == settings.SUPPORT_EMBEDDING_MODEL
    chunk_vectors = (
        [chunk.embedding for chunk in chunks] if current_model
        else await embed_texts([chunk.content for chunk in chunks])
    )
    reviewed_images = [image for image in images if image.description.strip()]
    image_vectors = (
        [image.embedding for image in reviewed_images if image.embedding is not None]
        if current_model and all(image.embedding is not None for image in reviewed_images)
        else await embed_texts([image.description for image in reviewed_images]) if reviewed_images else []
    )
    records = [
        SupportVector(
            source_id=document.id, item_id=chunk.id, source_type="document",
            revision=revision, title=document.title, content=chunk.content,
            url=(f"/api/v1/support/documents/{document.id}?page={chunk.page_number}" if chunk.page_number
                 else document.source_url or f"/api/v1/support/documents/{document.id}"),
            embedding=vector, page_number=chunk.page_number,
        )
        for chunk, vector in zip(chunks, chunk_vectors, strict=True)
    ]
    records.extend(
        SupportVector(
            source_id=document.id, item_id=image.id, source_type="image",
            revision=revision, title=f"{document.title} · image: {image.filename}", content=image.description,
            url=f"/api/v1/support/documents/{document.id}" + (f"?page={image.page_number}" if image.page_number else ""),
            embedding=vector, page_number=image.page_number, image_id=image.id,
        )
        for image, vector in zip(reviewed_images, image_vectors, strict=True)
    )
    return records


async def sync_published_document(document_id: str) -> bool:
    """Backfill one approved document into Milvus without changing its review state."""
    async with session_factory() as db:
        document = await db.get(SupportDocument, document_id)
        if document is None or document.status != "published":
            return False
        chunks = await db.exec(select(SupportChunk).where(col(SupportChunk.document_id) == document_id))
        images = await db.exec(select(SupportDocumentImage).where(col(SupportDocumentImage.document_id) == document_id))
        records = await document_vectors(document, list(chunks.all()), list(images.all()), document.updated_at.isoformat())
    await replace_source(document_id, records)
    return True


async def search_support_vectors(query: str, limit: int = 8) -> list[RetrievedPassage]:
    """Search all source types, then check live review state and revision in PostgreSQL."""
    if not query.strip() or not settings.SUPPORT_EMBEDDING_API_KEY:
        return []
    try:
        vector = (await embed_texts([query]))[0]
        hits = await search_milvus(vector, max(limit * 3, 18))
    except Exception:
        logger.exception("support_milvus_search_failed")
        return []
    article_ids = {str(hit["entity"]["source_id"]) for hit in hits if hit["entity"]["source_type"] == "article"}
    document_ids = {str(hit["entity"]["source_id"]) for hit in hits if hit["entity"]["source_type"] != "article"}
    async with session_factory() as db:
        articles = {entry.id: entry for entry in (await db.exec(
            select(SupportKnowledge).where(col(SupportKnowledge.id).in_(article_ids))
        )).all()} if article_ids else {}
        documents = {entry.id: entry for entry in (await db.exec(
            select(SupportDocument).where(col(SupportDocument.id).in_(document_ids))
        )).all()} if document_ids else {}
    passages: list[RetrievedPassage] = []
    for hit in hits:
        entity = hit["entity"]
        source_id = str(entity["source_id"])
        source_type = str(entity["source_type"])
        revision = str(entity["revision"])
        if source_type == "article":
            article = articles.get(source_id)
            if article is None or article.status != "published" or str(article.version) != revision:
                continue
        else:
            document = documents.get(source_id)
            if document is None or document.status != "published" or document.updated_at.isoformat() != revision:
                continue
        score = float(hit["distance"])
        if score < 0.28:
            continue
        page = int(entity["page_number"])
        passages.append(RetrievedPassage(
            id=str(hit["id"]), title=str(entity["title"]), content=str(entity["content"]),
            url=str(entity["url"]), distance=1 - score,
            page_number=page if page >= 0 else None,
            image_id=str(entity["image_id"]) or None,
        ))
        if len(passages) >= limit:
            break
    return passages


async def search_documents(query: str, limit: int = 5) -> list[RetrievedPassage]:
    """Fuse exact term and vector rankings over published passages."""
    if not query.strip():
        return []
    lexical_rows: list[tuple[SupportChunk, SupportDocument, float]] = []
    try:
        terms = list(dict.fromkeys(re.findall(r"[a-záéíóúüñ][a-záéíóúüñ0-9_]{2,}", query.casefold())))
        for phrase in re.findall(r"[\u3400-\u9fff]{2,}", query):
            terms.extend(phrase[start : start + 2] for start in range(max(0, len(phrase) - 1)))
        terms = list(dict.fromkeys(terms))[:12]
        async with session_factory() as db:
            available = await db.exec(
                select(SupportDocument.id)
                .where(
                    SupportDocument.status == "published",
                )
                .limit(1)
            )
            if available.first() is None:
                return []
            latin_terms = [term for term in terms if re.fullmatch(r"[a-záéíóúüñ][a-záéíóúüñ0-9_]{2,}", term)]
            keyword_filter = None
            keyword_score = literal_column("0")
            if terms:
                keyword_matches = [col(SupportChunk.keywords_json).ilike(f"%{term}%") for term in terms]
                keyword_filter = or_(*keyword_matches)
                keyword_score = sum((case((matched, 1), else_=0) for matched in keyword_matches), literal_column("0"))
            if terms:
                phrase = " | ".join(latin_terms)
                config = literal_column("'simple'::regconfig")
                text_vector = func.to_tsvector(config, col(SupportChunk.content))
                rank = (
                    func.ts_rank_cd(text_vector, func.to_tsquery(config, phrase))
                    if latin_terms
                    else literal_column("0")
                ) + keyword_score
                lexical_match = text_vector.op("@@")(func.to_tsquery(config, phrase)) if latin_terms else None
                if keyword_filter is not None:
                    lexical_match = keyword_filter if lexical_match is None else lexical_match | keyword_filter
                if lexical_match is None:
                    lexical_match = literal_column("FALSE")
                result = await db.exec(
                    select(SupportChunk, SupportDocument, rank.label("rank"))
                    .join(SupportDocument, col(SupportChunk.document_id) == col(SupportDocument.id))
                    .where(SupportDocument.status == "published", lexical_match)
                    .order_by(rank.desc())
                    .limit(12)
                )
                lexical_rows = [(chunk, document, float(score)) for chunk, document, score in result.all()]
    except Exception:
        logger.exception("support_lexical_search_failed")

    vector_rows: list[tuple[SupportChunk, SupportDocument, float]] = []
    if settings.SUPPORT_EMBEDDING_API_KEY:
        try:
            vector = (await embed_texts([query]))[0]
            distance = SupportChunk.embedding.cosine_distance(vector)  # pyright: ignore[reportAttributeAccessIssue]
            async with session_factory() as db:
                result = await db.exec(
                    select(SupportChunk, SupportDocument, distance.label("distance"))
                    .join(SupportDocument, col(SupportChunk.document_id) == col(SupportDocument.id))
                    .where(SupportDocument.status == "published")
                    .where(SupportDocument.embedding_model == settings.SUPPORT_EMBEDDING_MODEL)
                    .order_by(distance)
                    .limit(12)
                )
                vector_rows = [
                    (chunk, document, float(score))
                    for chunk, document, score in result.all()
                    if score is not None and float(score) <= 0.72
                ]
        except Exception:
            logger.exception("support_vector_search_failed")

    scored: dict[str, tuple[float, SupportChunk, SupportDocument, float]] = {}
    for rank, (chunk, document, _) in enumerate(lexical_rows, start=1):
        scored[chunk.id] = (1 / (60 + rank), chunk, document, 0.0)
    for rank, (chunk, document, distance_value) in enumerate(vector_rows, start=1):
        previous = scored.get(chunk.id)
        score = 1 / (60 + rank) + (previous[0] if previous else 0.0)
        scored[chunk.id] = (score, chunk, document, distance_value)
    ordered = sorted(scored.values(), key=lambda item: item[0], reverse=True)[:limit]
    return [
        RetrievedPassage(
            id=chunk.id,
            title=(f"{document.title} · v{document.game_version}" if document.game_version else document.title)
            + (f" · p. {chunk.page_number}" if chunk.page_number else ""),
            content=chunk.content,
            url=(
                f"/api/v1/support/documents/{document.id}?page={chunk.page_number}"
                if chunk.page_number
                else document.source_url or f"/api/v1/support/documents/{document.id}"
            ),
            distance=distance_value,
            page_number=chunk.page_number,
        )
        for _, chunk, document, distance_value in ordered
    ]
